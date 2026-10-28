"""Private uploads reach the resident's scoped reader without filesystem tools."""

from __future__ import annotations

import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from langchain_core.messages import AIMessage, HumanMessage

from octop.infra.projects import resident_tasks
from tests.integration.test_project_task_artifacts import real_model as real_model
from tests.integration.test_project_task_files import text_provider, tool_result, upload
from tests.integration.test_project_task_images import image_bytes
from tests.integration.test_project_tasks_ws import connected
from tests.integration.test_resident_conversations import _base, _binding, _turn
from tests.integration.test_resident_images import resident_model
from tests.integration.test_resident_stream_delivery import subscribe
from tests.support.auth import create_user
from tests.support.harness import patch_resident_stream

PROOF = "resident-upload-proof"


def document(kind):
    output = io.BytesIO()
    if kind == "txt":
        return "notes.txt", "text/plain", PROOF.encode()
    if kind == "pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=100)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        contents = DecodedStreamObject()
        contents.set_data(f"BT /F1 12 Tf 10 50 Td ({PROOF}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(contents)
        writer.write(output)
        return "notes.pdf", "application/pdf", output.getvalue()
    if kind == "docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph(PROOF)
        doc.save(output)
        return (
            "notes.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            output.getvalue(),
        )
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append([PROOF])
    workbook.save(output)
    workbook.close()
    return (
        "notes.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        output.getvalue(),
    )


def read_message(attachment_id, call_id):
    return AIMessage(
        content="",
        tool_calls=[
            {
                "id": call_id,
                "name": "resident_read_attachment",
                "args": {"attachment_id": attachment_id},
            }
        ],
    )


def text_model(monkeypatch, server):
    resident = server.app_runtime.resident_conversations
    configuration = resident._configuration

    def select(owner_id, agent_id):
        return {**configuration(owner_id, agent_id), "model_ref": "text-only/plain-chat"}

    monkeypatch.setattr(resident, "_configuration", select)


@pytest.mark.parametrize("kind", ["txt", "pdf", "docx", "xlsx"])
async def test_file_only_text_model_reads_real_content_and_can_read_it_again(
    env_with_provider, monkeypatch, real_model, kind
):
    client, server, auth = env_with_provider
    await text_provider(client, auth)
    text_model(monkeypatch, server)
    binding = await _binding(client, auth)
    base = _base(binding)
    name, media_type, data = document(kind)
    attachment = await upload(client, auth, name, data, media_type)
    identifier = attachment["attachment_id"]
    observed = resident_model(
        real_model,
        monkeypatch,
        [
            read_message(identifier, "first-file-read"),
            AIMessage(content="Read the actual uploaded file."),
            read_message(identifier, "followup-file-read"),
            AIMessage(content="Read the same original file again."),
        ],
    )
    body = {"request_id": uuid4().hex, "attachment_ids": [identifier]}
    accepted = await client.post(base + "/messages", headers=auth, json=body)
    assert accepted.status_code == 202, accepted.text
    first = await _turn(client, auth, base, body["request_id"], "completed")
    assert first["prompt"] == "" and first["attachments"] == [attachment]
    extracted = json.loads(tool_result(observed[1], "first-file-read"))
    assert PROOF in extracted["text"] and extracted["attachment_id"] == identifier
    assert extracted["next_offset"] is None
    user = next(message for message in observed[0] if isinstance(message, HumanMessage))
    assert identifier in str(user.content)
    assert attachment["workspace_path"] not in str(user.content)
    assert all(block["type"] == "text" for block in user.content)
    repeated = await client.post(base + "/messages", headers=auth, json=body)
    assert repeated.status_code == 202 and repeated.json() == first
    followup = {"request_id": uuid4().hex, "prompt": "Read the original uploaded file again."}
    sent = await client.post(base + "/messages", headers=auth, json=followup)
    assert sent.status_code == 202, sent.text
    second = await _turn(client, auth, base, followup["request_id"], "completed")
    assert second["attachments"] == []
    assert PROOF in json.loads(tool_result(observed[3], "followup-file-read"))["text"]
    assert len(observed) == 4
    for suffix in ("conversation", "conversation/history"):
        response = await client.get(base + "/" + suffix, headers=auth)
        assert response.status_code == 200, response.text
        turns = (
            response.json()["turns"]
            if suffix == "conversation"
            else [item["turn"] for item in response.json()["items"]]
        )
        assert turns == [first, second]
    if kind == "txt":
        session, ws = await connected(client, auth)
        try:
            snapshot = await subscribe(ws, binding, "private-file-snapshot")
            assert snapshot["turns"] == [first, second]
        finally:
            await session.__aexit__(None, None, None)
    owner = server.services.agent_repo.get(binding["agent_id"]).user_id
    receipt = server.services.resident_request_repo.get(owner, body["request_id"])
    configuration = json.loads(receipt.configuration_json)
    assert configuration["model_ref"] == "text-only/plain-chat"
    assert configuration["attachment_read_enabled"] is True
    # Uploading a private file does not implicitly authorize a delegated Task to read it.
    assert configuration["resident_task_context"]["message"]["attachment_ids"] == []
    workspace = Path(
        json.loads(server.services.agent_repo.get(binding["agent_id"]).config_json)["workspace_dir"]
    )
    assert not (workspace / attachment["workspace_path"]).exists()
    assert (await client.get(attachment["preview_url"], headers=auth)).content == data
    internal_id = resident_tasks.task_id(owner, binding["thread_id"])
    assert (await client.get("/api/project-tasks/" + internal_id, headers=auth)).status_code == 404


async def test_file_uploads_remain_owner_scoped_and_mixed_images_still_require_vision(
    env_with_provider, monkeypatch
):
    client, server, auth = env_with_provider
    binding = await _binding(client, auth)
    base = _base(binding)
    owner = server.services.agent_repo.get(binding["agent_id"]).user_id
    file = await upload(client, auth)
    image = await upload(client, auth, "image.png", image_bytes(), "image/png")
    foreign = await create_user(client, auth, username="private-file-other")
    other_file = await upload(client, foreign)
    denied = await client.post(
        base + "/messages",
        headers=auth,
        json={"request_id": uuid4().hex, "attachment_ids": [other_file["attachment_id"]]},
    )
    assert denied.status_code == 404, denied.text
    assert (await client.get(file["preview_url"], headers=foreign)).status_code == 404
    body = {
        "request_id": uuid4().hex,
        "attachment_ids": [file["attachment_id"], image["attachment_id"]],
    }
    with monkeypatch.context() as limited:
        limited.setattr(
            server.app_runtime.agent_registry.providers,
            "is_model_ref_multimodal",
            lambda _ref: False,
        )
        denied = await client.post(base + "/messages", headers=auth, json=body)
    assert (
        denied.status_code == 422 and denied.json()["error"]["code"] == "TASK_IMAGE_MODEL_REQUIRED"
    ), denied.text
    assert server.services.resident_request_repo.list(binding["thread_id"], owner) == []
    with server.services.db.read_snapshot() as conn:
        assert all(
            server.services.project_task_attachment_repo.get(conn, owner, identifier).task_id
            is None
            for identifier in body["attachment_ids"]
        )

    async def stream(_message):
        yield {"type": "token", "content": "Received both inputs."}

    patch_resident_stream(monkeypatch, stream)
    accepted = await client.post(base + "/messages", headers=auth, json=body)
    assert accepted.status_code == 202, accepted.text
    completed = await _turn(client, auth, base, body["request_id"], "completed")
    assert completed["attachments"] == [file, image]


async def test_legacy_gateway_rejects_new_files_and_text_followups_with_retained_files(
    env_with_provider, monkeypatch
):
    client, server, auth = env_with_provider
    binding = await _binding(client, auth)
    base = _base(binding)
    resident = server.app_runtime.resident_conversations
    await resident.close()
    attachment = await upload(client, auth)
    body = {"request_id": uuid4().hex, "attachment_ids": [attachment["attachment_id"]]}
    with monkeypatch.context() as legacy:
        legacy.setattr(resident.gateway, "_task_worker", None)
        denied = await client.post(base + "/messages", headers=auth, json=body)
    assert denied.status_code == 422 and denied.json()["error"]["code"] == "RESIDENT_TEXT_ONLY", (
        denied.text
    )
    accepted = await client.post(base + "/messages", headers=auth, json=body)
    assert accepted.status_code == 202, accepted.text
    with monkeypatch.context() as legacy:
        legacy.setattr(resident.gateway, "_task_worker", None)
        denied = await client.post(
            base + "/messages",
            headers=auth,
            json={"request_id": uuid4().hex, "prompt": "Read the previous file."},
        )
    assert denied.status_code == 422 and denied.json()["error"]["code"] == "RESIDENT_TEXT_ONLY", (
        denied.text
    )
