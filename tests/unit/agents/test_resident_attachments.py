"""Private attachment reads are bounded and cannot select host or other task files."""

from __future__ import annotations

import io
import json
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

import pytest
from tests.unit.agents.test_resident_profiles import profile_env as profile_env
from tests.unit.agents.test_resident_tools import tool_env as tool_env

from octop.infra.agents import resident_attachments, resident_tools
from octop.infra.agents.resident_identity import (
    ResidentTextOnlyMiddleware,
    resident_execution_scope,
)
from octop.infra.agents.resident_tools import resident_tools_scope
from octop.infra.db.repos.project_tasks import ProjectTaskRow
from octop.infra.projects import input_attachments, resident_tasks
from octop.infra.projects.tasks import now_ms


@pytest.fixture
async def attachment_env(tool_env):
    services, profile, source = tool_env
    identifier = await resident_tasks.ensure_task(services, 1, profile.thread_id, source.request_id)
    return services, profile, source, identifier


def _upload(env, name="notes.txt", data=b"Actual uploaded evidence", *, owner=1, task=None):
    services, _, _, identifier = env
    item = input_attachments.save(services, owner, name, data)
    with services.db.transaction() as conn:
        assert services.project_task_repo.lock_owner(conn, owner)
        input_attachments.bind(
            services, conn, owner, task or identifier, [item.attachment_id], resident=True
        )
    return item


@contextmanager
def _scope(env, *, enabled=True, turn=None):
    services, profile, source, _ = env
    with (
        resident_execution_scope(1, profile.thread_id),
        resident_tools_scope(
            services=services,
            owner_id=1,
            agent_id=profile.agent_id,
            thread_id=profile.thread_id,
            turn_id=turn or source.turn_id,
            attachment_read_enabled=enabled,
        ),
    ):
        yield


def _call(env, **arguments):
    _, profile, _, _ = env
    return resident_tools.execute_resident_tool(
        owner_id=1,
        thread_id=profile.thread_id,
        locale="en",
        tool_call={"name": "resident_read_attachment", "args": arguments, "id": uuid4().hex},
    )


def _json(message):
    return json.loads(message.content)


def _document(kind):
    buffer = io.BytesIO()
    if kind == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Actual Word evidence")
        document.save(buffer)
    elif kind == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Actual spreadsheet evidence", 42])
        workbook.save(buffer)
        workbook.close()
    elif kind == "pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (Actual PDF evidence) Tj ET")
        page[NameObject("/Contents")] = stream
        writer.write(buffer)
    elif kind == "pptx":
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Actual slide evidence"
        presentation.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("extension", "expected"),
    [
        ("docx", "Actual Word evidence"),
        ("xlsx", "Actual spreadsheet evidence"),
        ("pdf", "Actual PDF evidence"),
        ("pptx", "Actual slide evidence"),
    ],
)
async def test_existing_document_parsers_read_verified_uploads(attachment_env, extension, expected):
    item = _upload(attachment_env, f"report.{extension}", _document(extension))
    with _scope(attachment_env):
        result = _call(attachment_env, attachment_id=item.attachment_id)
    assert result.status == "success"
    assert expected in _json(result)["text"]
    assert "base64" not in result.content
    assert _json(result)["next_offset"] is None


async def test_text_pages_are_actual_character_windows_and_temp_copies_are_removed(
    attachment_env, monkeypatch
):
    text = "附件证据\n" * 8000
    item = _upload(attachment_env, data=text.encode())
    paths = []
    original = resident_attachments.parse_document

    def parse(path, *, ocr):
        assert ocr is None
        paths.append(path)
        assert "octop-resident-attachment-" in str(path)
        return original(path, ocr=ocr)

    monkeypatch.setattr(resident_attachments, "parse_document", parse)
    with _scope(attachment_env):
        first = _json(_call(attachment_env, attachment_id=item.attachment_id))
        second = _json(
            _call(attachment_env, attachment_id=item.attachment_id, offset=12000, limit=16000)
        )
        last = _json(_call(attachment_env, attachment_id=item.attachment_id, offset=len(text)))
    assert first["text"] == text[:12000] and first["next_offset"] == 12000
    assert second["text"] == text[12000:28000] and second["next_offset"] == 28000
    assert first["total_characters"] == len(text)
    assert last["text"] == "" and last["next_offset"] is None
    assert paths and all(not path.exists() for path in paths)


async def test_listing_is_metadata_only_paged_and_excludes_images(attachment_env, monkeypatch):
    first = _upload(attachment_env, name="first.txt")
    second = _upload(attachment_env, name="second.txt")
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (1, 1)).save(buffer, format="PNG")
    _upload(attachment_env, name="image.png", data=buffer.getvalue())
    download = Mock(side_effect=AssertionError("Listing must not load attachment bytes"))
    monkeypatch.setattr(input_attachments, "download", download)
    with _scope(attachment_env):
        all_files = _json(_call(attachment_env))
        page = _json(_call(attachment_env, offset=1))
        budget = len(json.dumps(all_files["items"][:1], ensure_ascii=False))
        bounded = _json(_call(attachment_env, limit=budget))
    assert [item["attachment_id"] for item in all_files["items"]] == [
        first.attachment_id,
        second.attachment_id,
    ]
    assert [item["attachment_id"] for item in page["items"]] == [second.attachment_id]
    assert all_files["total_items"] == 2 and page["next_offset"] is None
    assert bounded["items"] == all_files["items"][:1] and bounded["next_offset"] == 1
    assert len(json.dumps(bounded["items"], ensure_ascii=False)) <= budget
    assert "workspace_path" not in json.dumps(all_files)
    download.assert_not_called()


async def test_list_budget_too_small_fails_instead_of_returning_same_offset(attachment_env):
    _upload(attachment_env)
    with _scope(attachment_env):
        result = _call(attachment_env, limit=1)
    assert result.status == "error"
    assert _json(result)["error_code"] == "RESIDENT_TOOL_INPUT_INVALID"


@pytest.mark.parametrize("name", ["code.py", "events.log", "client.ts", "README"])
async def test_source_and_extensionless_utf8_text_are_readable(attachment_env, name):
    text = "# actual UTF-8 source\nprint('附件证据')\n"
    item = _upload(attachment_env, name=name, data=b"\xef\xbb\xbf" + text.encode())
    with _scope(attachment_env):
        result = _call(attachment_env, attachment_id=item.attachment_id)
    assert result.status == "success" and _json(result)["text"] == text


@pytest.mark.parametrize(
    "invalid",
    ["foreign_task", "foreign_owner", "pending", "host_path", "corrupted", "ordinary_task"],
)
async def test_reader_rejects_unbound_other_conversation_and_host_paths(
    attachment_env, monkeypatch, invalid
):
    services, _, _, _ = attachment_env
    if invalid in {"foreign_task", "foreign_owner"}:
        owner = 2 if invalid == "foreign_owner" else 1
        identifier = "other-hidden-task"
        with services.db.transaction() as conn:
            services.project_task_repo.insert(
                conn,
                ProjectTaskRow(
                    identifier,
                    owner,
                    None,
                    "Other",
                    "other",
                    None,
                    now_ms(),
                    now_ms(),
                    internal_kind="resident",
                ),
            )
        item = _upload(attachment_env, owner=owner, task=identifier)
        attachment_id = item.attachment_id
    elif invalid == "host_path":
        attachment_id = str(Path("private") / "system.txt")
    elif invalid == "pending":
        item = input_attachments.save(services, 1, "pending.txt", b"Not yet selected")
        attachment_id = item.attachment_id
    else:
        item = _upload(attachment_env)
        attachment_id = item.attachment_id
        with services.db.transaction() as conn:
            if invalid == "ordinary_task":
                conn.execute(
                    "UPDATE project_tasks SET internal_kind='' WHERE task_id=?",
                    (attachment_env[3],),
                )
            else:
                conn.execute(
                    "UPDATE project_task_attachments SET content=? WHERE attachment_id=?",
                    (b"changed", attachment_id),
                )
    parse = Mock(side_effect=AssertionError("Unauthorized files must not be parsed"))
    monkeypatch.setattr(resident_attachments, "parse_document", parse)
    with _scope(attachment_env):
        result = _call(attachment_env, attachment_id=attachment_id)
    assert result.status == "error" and _json(result)["error_code"] == "TASK_ATTACHMENT_UNAVAILABLE"
    parse.assert_not_called()


@pytest.mark.parametrize(
    "invalid",
    [
        "old_scope",
        "cancelled",
        "expired",
        "wrong_turn",
        "path_argument",
        "negative_offset",
        "over_limit",
    ],
)
async def test_tool_respects_frozen_scope_live_request_and_arguments(attachment_env, invalid):
    services, profile, source, _ = attachment_env
    item = _upload(attachment_env)
    arguments = {"attachment_id": item.attachment_id}
    if invalid == "cancelled":
        services.resident_request_repo.cancel(
            1, profile.thread_id, "cancel-read", source.turn_id, now_ms()
        )
    elif invalid == "expired":
        assert services.resident_request_repo.renew(
            source.turn_id, "worker", now_ms(), now_ms() - 1
        )
    elif invalid == "path_argument":
        arguments["path"] = "system.txt"
    elif invalid == "negative_offset":
        arguments["offset"] = -1
    elif invalid == "over_limit":
        arguments["limit"] = 16001
    with _scope(
        attachment_env,
        enabled=invalid != "old_scope",
        turn="wrong" if invalid == "wrong_turn" else None,
    ):
        result = _call(attachment_env, **arguments)
    assert result.status == "error"
    expected = (
        "RESIDENT_TOOL_INPUT_INVALID"
        if invalid in {"path_argument", "negative_offset", "over_limit"}
        else "FORBIDDEN"
    )
    assert _json(result)["error_code"] == expected


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("archive.zip", b"PK\x03\x04not opened"),
        ("empty.txt", b""),
        ("broken.docx", b"broken"),
        ("broken.pdf", b"plain UTF-8 is not a PDF"),
        ("raw.bin", b"\xff\xfe\x01"),
        ("unknown", b"NUL\x00byte"),
        ("control.log", b"binary\x01control"),
    ],
)
async def test_unreadable_formats_return_localized_tool_error(attachment_env, name, data):
    item = _upload(attachment_env, name=name, data=data)
    with _scope(attachment_env):
        result = _call(attachment_env, attachment_id=item.attachment_id)
    assert result.status == "error"
    assert _json(result)["error_code"] == "RESIDENT_ATTACHMENT_UNREADABLE"
    assert "message" in _json(result)


async def test_attachment_capability_adds_only_dedicated_reader_and_never_system_content(
    attachment_env,
):
    _, profile, _, _ = attachment_env
    item = _upload(attachment_env, data=b"PRIVATE FILE CONTENT")
    guard = ResidentTextOnlyMiddleware("en", 1, profile.thread_id)
    request = SimpleNamespace(override=lambda **kwargs: kwargs)
    with _scope(attachment_env):
        prepared = guard.wrap_model_call(request, lambda update: update)
    assert {tool["function"]["name"] for tool in prepared["tools"]} == {"resident_read_attachment"}
    assert "PRIVATE FILE CONTENT" not in prepared["system_message"].content
    with _scope(attachment_env, enabled=False):
        assert guard.wrap_model_call(request, lambda update: update)["tools"] == []
    assert item.attachment_id
