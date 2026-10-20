"""Project/task identity and labels in token usage exports."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from octop.api.common.usage_xlsx import build_usage_xlsx
from octop.infra.db.repos.usage import UsageRow


@pytest.mark.parametrize(
    ("locale", "project_sheet", "task_sheet", "no_project"),
    [("zh", "按项目", "按任务", "未关联项目"), ("en", "By project", "By task", "No project")],
)
def test_export_project_task_identity_is_literal_text(
    locale: str, project_sheet: str, task_sheet: str, no_project: str
) -> None:
    rows = [
        UsageRow(
            id=index,
            ts=1_700_000_000 + index,
            agent_id="",
            user_id=1,
            thread_id=f"task-{index}",
            model="provider:model",
            input_tokens=10,
            uncached_input_tokens=0,
            cache_read_tokens=0,
            cache_write_tokens=0,
            output_tokens=5,
            reasoning_tokens=0,
            total_tokens=15,
            model_calls=0,
            source="project_task",
            project_id="project-1" if index == 1 else "",
            project_name="=project" if index == 1 else "",
            task_id=f"task-{index}",
            task_title=f"=task-{index}",
        )
        for index in (1, 2)
    ]
    metrics = {"input_tokens": 10, "output_tokens": 5, "total_tokens": 15, "turns": 1}
    payload = build_usage_xlsx(
        rows=rows,
        by_day=[],
        by_agent=[],
        by_model=[],
        by_project=[
            {"key": "project-1", "label": "=project", **metrics},
            {"key": "", "label": "", **metrics},
        ],
        by_task=[{"key": row.task_id, "label": row.task_title, **metrics} for row in rows],
        agent_names={},
        usernames={1: "owner"},
        timezone="Asia/Shanghai",
        locale=locale,
    )
    workbook = load_workbook(BytesIO(payload))
    detail = workbook.worksheets[0]
    assert detail.auto_filter.ref == "A1:T3"
    assert [detail.cell(2, col).value for col in range(17, 21)] == [
        "project-1",
        "=project",
        "task-1",
        "=task-1",
    ]
    assert detail.cell(3, 18).value == no_project
    assert detail.cell(2, 18).data_type == "s"
    assert detail.cell(2, 20).data_type == "s"
    assert detail.cell(5, 14).value == "=SUM(N2:N3)"
    assert workbook[project_sheet].cell(2, 1).value == "=project"
    assert workbook[project_sheet].cell(2, 1).data_type == "s"
    assert workbook[project_sheet].cell(3, 1).value == no_project
    assert workbook[task_sheet].cell(2, 1).value == "=task-1"
    assert workbook[task_sheet].cell(2, 1).data_type == "s"
    assert workbook[project_sheet]._charts
    assert workbook[task_sheet]._charts
